"""Parse Excel registry files from the Z-Wave specs."""

from __future__ import annotations

import logging
import re
from pathlib import Path

import openpyxl

from mcp_zwave_specs.models import RegistryData

logger = logging.getLogger(__name__)

# Registry files mapped to their data sheet names and human-readable names
REGISTRY_FILES: dict[str, tuple[str, str]] = {
    # filename → (sheet_name, display_name)
    "Notification Command Class, list of assigned Notifications.xlsx": (
        "Notifications",
        "Notification Types",
    ),
    "Multilevel Sensor Command Class, list of assigned Multilevel Sensor types and scales.xlsx": (
        "Multilevel Sensor",
        "Sensor Types",
    ),
    "Indicator Command Class, list of assigned indicators and Property IDs.xlsx": (
        "Indicators",
        "Indicator Types",
    ),
    "Meter Table Monitor Command Class, list of assigned types, scales and datasets.xlsx": (
        "Meter Table",
        "Meter Table Types",
    ),
    "SDS14622 Anti-Theft Command Class, list of assigned Locking Entity IDs.xlsx": (
        "Locking Entity IDs",
        "Anti-Theft Locking Entities",
    ),
    "Simple AV Command Class, list of assigned AV Control codes.xlsx": (
        "Simple AV Codes",
        "AV Control Codes",
    ),
    "Z-Wave Manufacturer ID List.xlsx": (
        "Manufacturer ID",
        "Manufacturers",
    ),
    "Z-Wave Plus Assigned Icon Types.xlsx": (
        "Icon Type",
        "Icon Types",
    ),
    "Association Command Class, list of mandatory commands"
    " for the Lifeline Association Group.xlsx": (
        "Lifeline Reports",
        "Lifeline Association Commands",
    ),
}

# The CC list is in a different directory
CC_LIST_FILE = "List of defined Z-Wave Command Classes.xlsx"
CC_LIST_SHEETS = {
    "Command Class": "Command Classes",
    "Commands": "CC Commands",
}


def _find_header_row(ws: openpyxl.worksheet.worksheet.Worksheet) -> int:
    """Find the first row with multiple non-None cells (likely the header)."""
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        non_none = [c for c in row if c is not None]
        if len(non_none) >= 2:
            return i
    return 1


def _clean_cell(value: object) -> str:
    """Convert a cell value to a clean string."""
    if value is None:
        return ""
    s = str(value).strip()
    # Strip Excel HYPERLINK formulas to just the display text
    m = re.match(r'=HYPERLINK\([^,]+,"([^"]+)"\)', s)
    if m:
        return m.group(1)
    return s


def _parse_sheet(
    ws: openpyxl.worksheet.worksheet.Worksheet,
) -> tuple[list[str], list[dict[str, str]]]:
    """Parse a worksheet into column headers and row dicts."""
    header_row_num = _find_header_row(ws)

    rows_iter = ws.iter_rows(values_only=True)
    # Skip to header row
    headers_raw = None
    for i, row in enumerate(rows_iter, start=1):
        if i == header_row_num:
            headers_raw = row
            break

    if not headers_raw:
        return [], []

    headers = [_clean_cell(h) or f"col_{i}" for i, h in enumerate(headers_raw)]

    rows: list[dict[str, str]] = []
    for row in rows_iter:
        values = [_clean_cell(c) for c in row]
        # Skip completely empty rows
        if not any(values):
            continue
        row_dict = dict(zip(headers, values, strict=False))
        rows.append(row_dict)

    return headers, rows


def parse_registries(registries_dir: Path) -> dict[str, RegistryData]:
    """Parse all registry Excel files."""
    result: dict[str, RegistryData] = {}

    for filename, (sheet_name, display_name) in REGISTRY_FILES.items():
        filepath = registries_dir / filename
        if not filepath.exists():
            logger.warning("Registry file not found: %s", filename)
            continue

        try:
            wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
        except Exception:
            logger.warning("Failed to open %s", filename, exc_info=True)
            continue

        # Try exact sheet name, then case-insensitive match
        ws = None
        if sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
        else:
            for name in wb.sheetnames:
                if name.lower() == sheet_name.lower():
                    ws = wb[name]
                    break

        if ws is None:
            # Fall back to first non-metadata sheet
            for name in wb.sheetnames:
                if name.lower() not in ("frontpage", "front page", "changes", "change log"):
                    ws = wb[name]
                    break

        if ws is None:
            logger.warning("No data sheet found in %s", filename)
            wb.close()
            continue

        columns, rows = _parse_sheet(ws)
        wb.close()

        # Carry forward empty first-column values (e.g., Lifeline sheet groups
        # multiple commands under one CC name)
        if rows and columns:
            first_col = columns[0]
            prev_val = ""
            for row in rows:
                if row.get(first_col):
                    prev_val = row[first_col]
                else:
                    row[first_col] = prev_val

        key = display_name.lower().replace(" ", "_")
        result[key] = RegistryData(
            name=display_name,
            filename=filename,
            rows=rows,
            columns=columns,
        )
        logger.info("Parsed %d rows from %s", len(rows), display_name)

    return result


def parse_cc_list(cc_list_path: Path) -> dict[str, RegistryData]:
    """Parse the master Command Class list Excel file."""
    result: dict[str, RegistryData] = {}

    if not cc_list_path.exists():
        logger.warning("CC list file not found: %s", cc_list_path)
        return result

    try:
        wb = openpyxl.load_workbook(cc_list_path, read_only=True, data_only=True)
    except Exception:
        logger.warning("Failed to open CC list", exc_info=True)
        return result

    for sheet_name, display_name in CC_LIST_SHEETS.items():
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        columns, rows = _parse_sheet(ws)
        key = display_name.lower().replace(" ", "_")
        result[key] = RegistryData(
            name=display_name,
            filename=CC_LIST_FILE,
            rows=rows,
            columns=columns,
        )
        logger.info("Parsed %d rows from %s", len(rows), display_name)

    wb.close()
    return result
